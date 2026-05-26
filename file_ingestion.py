"""
file_ingestion.py
=================
Extracts plain text from uploaded files so the content can be stored in the
DB and injected into the LLM prompt as context.

Supported types
---------------
Plain text  : .txt .md .py .js .ts .jsx .tsx .css .html .cs .cpp .c .java
              .json .yaml .yml .sh .sql .xml .toml .ini .env  (built-in)
CSV         : .csv          (built-in csv module)
PDF         : .pdf          (pypdf  — pip install pypdf)
Word        : .docx         (python-docx — already installed)
Excel       : .xlsx         (openpyxl   — already installed)

Usage
-----
    from file_ingestion import extract_text, ExtractionError

    try:
        text = extract_text(file_bytes, ".pdf")
    except ExtractionError as e:
        # unsupported type or parse failure
        ...
"""

from __future__ import annotations

import csv
import io

# ── Optional library imports ───────────────────────────────────────────────────
try:
    from pypdf import PdfReader
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False


# ── Constants ──────────────────────────────────────────────────────────────────
PLAIN_TEXT_EXTENSIONS = {
    ".txt", ".md", ".py", ".js", ".ts", ".jsx", ".tsx",
    ".css", ".html", ".htm", ".cs", ".cpp", ".c", ".h",
    ".java", ".json", ".yaml", ".yml", ".sh", ".sql",
    ".xml", ".toml", ".ini", ".env",
}

# All extensions this module can handle
SUPPORTED_EXTENSIONS = PLAIN_TEXT_EXTENSIONS | {".csv", ".pdf", ".docx", ".xlsx"}

# Soft cap — truncate extracted text to this many characters before storing.
# Keeps very large files from blowing out the context window.
MAX_CHARS = 40_000


class ExtractionError(Exception):
    """Raised when a file cannot be parsed."""


# ── Public API ─────────────────────────────────────────────────────────────────
def extract_text(file_bytes: bytes, extension: str) -> str:
    """
    Extract and return plain text from *file_bytes*.

    Parameters
    ----------
    file_bytes : raw bytes of the uploaded file
    extension  : lower-cased file extension including the dot, e.g. ".pdf"

    Returns
    -------
    Extracted text, truncated to MAX_CHARS if necessary.

    Raises
    ------
    ExtractionError  if the type is unsupported or parsing fails.
    """
    ext = extension.lower()

    if ext in PLAIN_TEXT_EXTENSIONS:
        text = _extract_plain(file_bytes)
    elif ext == ".csv":
        text = _extract_csv(file_bytes)
    elif ext == ".pdf":
        text = _extract_pdf(file_bytes)
    elif ext == ".docx":
        text = _extract_docx(file_bytes)
    elif ext == ".xlsx":
        text = _extract_xlsx(file_bytes)
    else:
        raise ExtractionError(f"Unsupported file type: '{ext}'")

    text = text.strip()
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + f"\n\n[… truncated at {MAX_CHARS:,} characters]"

    return text


def is_supported(extension: str) -> bool:
    """Return True if this module can extract text from the given extension."""
    return extension.lower() in SUPPORTED_EXTENSIONS


# ── Extractors ─────────────────────────────────────────────────────────────────
def _extract_plain(data: bytes) -> str:
    """Decode bytes as UTF-8, falling back to latin-1."""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _extract_csv(data: bytes) -> str:
    """Convert CSV rows to a readable pipe-delimited table."""
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")

    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return ""

    # Format as aligned columns
    col_widths = [max(len(str(row[i])) for row in rows if i < len(row))
                  for i in range(max(len(r) for r in rows))]

    lines = []
    for idx, row in enumerate(rows):
        padded = " | ".join(str(cell).ljust(col_widths[i]) for i, cell in enumerate(row))
        lines.append(padded)
        if idx == 0:  # separator under header
            lines.append("-+-".join("-" * w for w in col_widths))

    return "\n".join(lines)


def _extract_pdf(data: bytes) -> str:
    if not PDF_AVAILABLE:
        raise ExtractionError("PDF support requires 'pypdf': pip install pypdf")
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(f"[Page {i}]\n{page_text}")
        return "\n\n".join(pages)
    except Exception as exc:
        raise ExtractionError(f"PDF parse error: {exc}") from exc


def _extract_docx(data: bytes) -> str:
    if not DOCX_AVAILABLE:
        raise ExtractionError("DOCX support requires 'python-docx': pip install python-docx")
    try:
        doc = DocxDocument(io.BytesIO(data))
        return "\n".join(para.text for para in doc.paragraphs)
    except Exception as exc:
        raise ExtractionError(f"DOCX parse error: {exc}") from exc


def _extract_xlsx(data: bytes) -> str:
    if not XLSX_AVAILABLE:
        raise ExtractionError("XLSX support requires 'openpyxl': pip install openpyxl")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    rows.append("\t".join("" if c is None else str(c) for c in row))
            if rows:
                sections.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
        return "\n\n".join(sections)
    except Exception as exc:
        raise ExtractionError(f"XLSX parse error: {exc}") from exc
